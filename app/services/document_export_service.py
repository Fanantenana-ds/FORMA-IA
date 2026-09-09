import io
from docx import Document as DocxDocument
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm

from app.models.document import Document, TypeDocument, FormatExport

class DocumentExportService:
    def exporter(self, document: Document, format: FormatExport) -> tuple[bytes, str, str]:
        """Retourne contenu binaire, nom de fichier, media_type"""
        if format == FormatExport.DOCX:
            return self._exporter_word(document)

        if format == FormatExport.PDF:
            return self._exporter_pdf(document)

        raise ValueError(f"Format d'export non supporté : {format}")


    def _titre_document(self, document: Document) -> str:
        titres = {
            TypeDocument.TDR: "Termes de Références",
            TypeDocument.OFFRE: "Offre Technique et Financière",
            TypeDocument.ATTESTATION: "Attestation de Formation",
        }
        return titres.get(document.type, "Document")


    def _ligne_contenu(self, document: Document) -> list[str]:
        if document.type == TypeDocument.TDR:
            return [
                f"Client : {document.client or '-'}",
                f"Objectifs : {document.objectifs or '-'}",
                f"Budget : {document.montant or '-'}",
            ]

        if document.type == TypeDocument.ATTESTATION:
            return [f"Numéro d'attestation : {document.numero_unique or '-'}"]
        return [document.contenu or ""]


    def _exporter_word(self, document: Document) -> tuple[bytes, str, str]:
        doc = DocxDocument()
        doc.add_heading(self._titre_document(document), level=1)

        for ligne in self._ligne_contenu(document):
            doc.add_paragraph(ligne)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        nom_fichier = f"{document.type.value.lower()}_{document.id}.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        return buffer.read(), nom_fichier, media_type


    def _exporter_pdf(self, document: Document) -> tuple[bytes, str, str]:
        buffer = io.BytesIO()
        c = canvas.Canvas(buffer, pagesize=A4)
        width, height = A4

        y = height - 3 *cm
        c.setFont("Helvetica-Bold", 16)
        c.drawString(2 * cm, y, self._titre_document(document))
        y -= 1.5 * cm

        c.setFont("Helvetica", 11)
        for ligne in self._ligne_contenu(document):
            for sous_ligne in self._decouper_ligne(ligne, 90):
                c.drawString(2 * cm, y, sous_ligne)
                y -= 0.7 * cm
                if y < 2 * cm:
                    c.showPage()
                    y = height - 2 * cm

        c.save()
        buffer.seek(0)

        nom_fichier = f"{document.type.value.lower()}_{document.id}.pdf"

        return buffer.read(), nom_fichier, "application/pdf"


    @staticmethod
    def _decouper_ligne(texte: str, largeur_max: int) -> list[str]:
        mots = texte.split(" ")
        lignes, ligne_courante = [], ""
        for mot in mots:
            if len(ligne_courante) + len(mot) + 1 <= largeur_max:
                ligne_courante = f"{ligne_courante} {mot}".strip()

            else:
                lignes.append(ligne_courante)
                ligne_courante = mot

        if ligne_courante:
            lignes.append(ligne_courante)

            return lignes or [""]