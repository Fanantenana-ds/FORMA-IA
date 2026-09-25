# app/services/rag/rag_document_generator.py
# ============================================================
# GÉNÉRATEUR DOCX — Portfolio et Syllabus (RAG, Étape G)
# ============================================================
# ÉCART ASSUMÉ ET DOCUMENTÉ par rapport à "même mécanisme que M2/TDR" :
# M2/TDR utilise docxtpl + un template .docx pré-conçu dans Word (mise en
# page soignée par un humain). Je ne peux pas concevoir ni vérifier
# visuellement un template Word avec mes outils texte. Construit donc
# directement avec python-docx (déjà une dépendance du projet) : le
# résultat est un vrai .docx structuré selon les gabarits de la mission,
# mais sans mise en forme graphique avancée (logo remplacé par un texte
# indicatif). Export PDF non inclus ici — comme demandé, uniquement après
# approbation HITL, réutilisable avec le même mécanisme docx2pdf/reportlab
# que tdr_document_generator.py si besoin.
# ============================================================

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

BASE_DIR = Path(__file__).resolve().parents[3]
EXPORTS_DIR = BASE_DIR / "exports" / "rag"


def _page_de_garde(doc: Document, titre: str, sous_titre: str, prepare_pour: str) -> None:
    doc.add_paragraph("[LOGO ALTIORA PREST]").alignment = WD_ALIGN_PARAGRAPH.CENTER
    titre_p = doc.add_heading(titre, level=0)
    titre_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous_p = doc.add_paragraph(sous_titre)
    sous_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Date : {datetime.now().strftime('%d/%m/%Y')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Préparé pour : {prepare_pour}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def _ajouter_tableau(doc: Document, entetes: list, lignes: list) -> None:
    table = doc.add_table(rows=1, cols=len(entetes))
    table.style = "Light Grid Accent 1"
    for i, entete in enumerate(entetes):
        table.rows[0].cells[i].text = str(entete)
    for ligne in lignes:
        cellules = table.add_row().cells
        for i, valeur in enumerate(ligne):
            cellules[i].text = str(valeur)


# ============================================================
# SYLLABUS
# ============================================================

def generer_docx_syllabus(donnees: Dict[str, Any], nom_fichier: str) -> Path:
    """Construit le DOCX du syllabus selon le gabarit de la mission (10
    sections + annexe interne des sources)."""
    doc = Document()

    _page_de_garde(
        doc, donnees.get("intitule", "[À compléter]"),
        f"Syllabus — {donnees.get('code', '')} — v{donnees.get('version', '1.0')}",
        donnees.get("prepare_pour", "[À compléter]"),
    )

    doc.add_heading("1. Fiche synthétique", level=1)
    _ajouter_tableau(doc, ["Champ", "Valeur"], [
        ["Intitulé", donnees.get("intitule", "[À compléter]")],
        ["Code", donnees.get("code", "[À compléter]")],
        ["Domaine", donnees.get("domaine", "[À compléter]")],
        ["Niveau", donnees.get("niveau", "[À compléter]")],
        ["Durée (jours)", donnees.get("duree_jours", "[À compléter]")],
        ["Durée (heures)", donnees.get("duree_heures", "[À compléter]")],
        ["Format", donnees.get("format", "[À compléter]")],
        ["Public cible", donnees.get("public_cible", "[À compléter]")],
        ["Prérequis", donnees.get("prerequis", "[À compléter]")],
        ["Effectif recommandé", donnees.get("effectif_recommande", "[À compléter]")],
        ["Langue", donnees.get("langue", "Français")],
        ["Lieu", donnees.get("lieu", "[À compléter]")],
    ])

    doc.add_heading("2. Contexte et enjeux", level=1)
    doc.add_paragraph(donnees.get("contexte_enjeux", "[À compléter]"))

    doc.add_heading("3. Objectifs pédagogiques", level=1)
    doc.add_paragraph(donnees.get("objectif_general", "[À compléter]"))
    for obj in donnees.get("objectifs_operationnels", []):
        doc.add_paragraph(obj, style="List Bullet")

    doc.add_heading("4. Compétences visées", level=1)
    _ajouter_tableau(
        doc, ["Compétence", "Niveau visé"],
        [[c.get("competence", ""), c.get("niveau_vise", "")] for c in donnees.get("competences", [])],
    )

    doc.add_heading("5. Programme détaillé", level=1)
    _ajouter_tableau(
        doc, ["Module", "Contenus", "Méthodes", "Durée", "Objectif lié"],
        [[m.get("titre", ""), m.get("contenus", ""), m.get("methodes", ""),
          m.get("duree", ""), m.get("objectif_lie", "")] for m in donnees.get("modules", [])],
    )
    if donnees.get("verification_durees_message"):
        p = doc.add_paragraph(f"⚠ {donnees['verification_durees_message']}")
        p.runs[0].italic = True

    doc.add_heading("6. Méthodes et moyens pédagogiques", level=1)
    doc.add_paragraph(donnees.get("methodes_pedagogiques", "[À compléter]"))

    doc.add_heading("7. Modalités d'évaluation", level=1)
    doc.add_paragraph(donnees.get("modalites_evaluation", "[À compléter]"))

    doc.add_heading("8. Supports remis aux participants", level=1)
    for support in donnees.get("supports_remis", []) or ["[À compléter]"]:
        doc.add_paragraph(support, style="List Bullet")

    doc.add_heading("9. Profil de l'intervenant", level=1)
    doc.add_paragraph(donnees.get("profil_intervenant", "[À compléter]"))

    doc.add_heading("10. Organisation pratique", level=1)
    doc.add_paragraph(donnees.get("organisation_pratique", "[À compléter] (horaires à confirmer)"))

    if donnees.get("sources_par_section"):
        doc.add_page_break()
        doc.add_heading("Annexe interne — Sources (usage HITL uniquement)", level=1)
        for section, sources in donnees["sources_par_section"].items():
            doc.add_paragraph(f"{section} : {', '.join(sources)}")

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    chemin = EXPORTS_DIR / nom_fichier
    doc.save(str(chemin))
    logger.info(f"📄 Syllabus DOCX généré : {chemin}")
    return chemin


# ============================================================
# PORTFOLIO
# ============================================================

def generer_docx_portfolio(donnees: Dict[str, Any], nom_fichier: str) -> Path:
    """Construit le DOCX du portfolio selon le gabarit de la mission (6
    sections + annexe interne des sources)."""
    doc = Document()

    _page_de_garde(
        doc, f"Références et expériences en formation — {donnees.get('domaine', '[À compléter]')}",
        f"Référence AO : {donnees.get('reference_ao', '[À compléter]')}", "",
    )

    doc.add_heading("1. Présentation d'ALTIORA PREST", level=1)
    doc.add_paragraph(donnees.get("presentation_altiora", "[À compléter par ALTIORA]"))

    doc.add_heading("2. Synthèse chiffrée", level=1)
    synthese = donnees.get("synthese_chiffree", {})
    _ajouter_tableau(doc, ["Indicateur", "Valeur"], [
        ["Formations réalisées", synthese.get("nb_formations_realisees", "[À compléter]")],
        ["Participants formés", synthese.get("nb_participants_total", "[À compléter]")],
        ["Clients distincts", synthese.get("nb_clients_distincts", "[À compléter]")],
        ["Période couverte", f"{synthese.get('periode_debut', '?')} — {synthese.get('periode_fin', '?')}"],
    ])

    doc.add_heading("3. Tableau récapitulatif", level=1)
    lignes = donnees.get("tableau_recapitulatif", [])
    _ajouter_tableau(
        doc, ["N°", "Intitulé", "Client", "Année", "Durée", "Participants", "Lieu"],
        [[i + 1, l["intitule"], l["client"], l["annee"], l["duree"], l["participants"], l["lieu"]]
         for i, l in enumerate(lignes)],
    )

    doc.add_heading("4. Fiches de référence", level=1)
    for fiche in donnees.get("fiches_reference", []):
        doc.add_heading(fiche.get("intitule", "[À compléter]"), level=2)
        _ajouter_tableau(doc, ["Champ", "Valeur"], [
            ["Client", fiche.get("client", "[À compléter]")],
            ["Période", fiche.get("periode", "[À compléter]")],
            ["Durée", fiche.get("duree", "[À compléter]")],
            ["Participants", fiche.get("participants", "[À compléter]")],
            ["Public", fiche.get("public", "[À compléter]")],
            ["Lieu", fiche.get("lieu", "[À compléter]")],
            ["Attestation de bonne exécution", fiche.get("attestation_disponible", "[À compléter]")],
        ])
        doc.add_paragraph(f"Objectifs : {fiche.get('objectifs', '[À compléter]')}")
        doc.add_paragraph("Contenus clés :")
        for contenu in fiche.get("contenus_cles", []) or ["[À compléter]"]:
            doc.add_paragraph(contenu, style="List Bullet")
        if fiche.get("resultats"):
            doc.add_paragraph(f"Résultats : {fiche['resultats'].get('note', '')}")

    doc.add_heading("5. Compétences et moyens mobilisés", level=1)
    doc.add_paragraph(donnees.get("competences_moyens", "[À compléter]"))

    doc.add_heading("6. Contact", level=1)
    doc.add_paragraph(donnees.get("contact", "[À compléter]"))

    if donnees.get("sources_par_fiche"):
        doc.add_page_break()
        doc.add_heading("Annexe interne — Sources (usage HITL uniquement)", level=1)
        for fiche_nom, sources in donnees["sources_par_fiche"].items():
            doc.add_paragraph(f"{fiche_nom} : {', '.join(sources)}")

    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    chemin = EXPORTS_DIR / nom_fichier
    doc.save(str(chemin))
    logger.info(f"📄 Portfolio DOCX généré : {chemin}")
    return chemin
