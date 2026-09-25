# ============================================================
# TESTS — app/services/rag/document_loader_service.py (Étape C §3)
# ============================================================
# Fichiers réels minimalistes créés à la volée (python-docx/pptx/openpyxl,
# reportlab pour le PDF) — aucun réseau, aucune base.
# ============================================================

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from app.services.rag.document_loader_service import (
    detect_file_type, extract_segments, pdf_necessite_ocr, compter_pages_pdf,
    Segment, FORMATS_IMAGE, FORMATS_INDEXABLES,
)


def test_detect_file_type_image_jpg(tmp_path):
    chemin = tmp_path / "photo.jpg"
    # Signature JPEG minimale reconnue par `filetype`
    chemin.write_bytes(bytes.fromhex("FFD8FFE000104A46494600010100000100010000FFD9"))
    assert detect_file_type(str(chemin)) in FORMATS_IMAGE


def test_detect_file_type_png(tmp_path):
    chemin = tmp_path / "photo.png"
    chemin.write_bytes(bytes.fromhex("89504E470D0A1A0A0000000D49484452"))
    assert detect_file_type(str(chemin)) in FORMATS_IMAGE


def test_detect_file_type_inconnu(tmp_path):
    chemin = tmp_path / "mystere.xyz"
    chemin.write_bytes(b"donnees quelconques")
    assert detect_file_type(str(chemin)) == "unknown"


def test_extract_docx_segment_unique(tmp_path):
    chemin = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_paragraph("Introduction au test.")
    doc.add_paragraph("Deuxième paragraphe.")
    doc.save(str(chemin))

    segments = extract_segments(str(chemin), "docx")

    assert len(segments) == 1  # DOCX : pas de pagination native
    assert segments[0].numero == 1
    assert "Introduction au test" in segments[0].texte


def test_extract_pptx_un_segment_par_diapositive(tmp_path):
    chemin = tmp_path / "presentation.pptx"
    prs = Presentation()
    layout = prs.slide_layouts[1]

    for i in range(3):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Titre diapositive {i + 1}"
        slide.placeholders[1].text = f"Contenu {i + 1}"

    prs.save(str(chemin))

    segments = extract_segments(str(chemin), "pptx")

    assert len(segments) == 3
    assert [s.numero for s in segments] == [1, 2, 3]
    assert "Titre diapositive 2" in segments[1].texte


def test_extract_xlsx_un_segment_par_feuille(tmp_path):
    chemin = tmp_path / "classeur.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Feuille1"
    ws1.append(["Nom", "Score"])
    ws1.append(["Alice", 90])
    ws2 = wb.create_sheet("Feuille2")
    ws2.append(["Autre", "Donnee"])
    wb.save(str(chemin))

    segments = extract_segments(str(chemin), "xlsx")

    assert len(segments) == 2
    assert segments[0].numero == 1
    assert segments[1].numero == 2
    assert "Alice" in segments[0].texte


def test_extract_pdf_texte_reel(tmp_path):
    chemin = tmp_path / "doc.pdf"
    c = canvas.Canvas(str(chemin))
    c.drawString(100, 750, "Ceci est une page de test avec du texte reel.")
    c.showPage()
    c.drawString(100, 750, "Deuxieme page du document de test.")
    c.showPage()
    c.save()

    segments = extract_segments(str(chemin), "pdf")
    nb_pages = compter_pages_pdf(str(chemin))

    assert nb_pages == 2
    assert len(segments) == 2
    assert segments[0].numero == 1
    assert segments[1].numero == 2


def test_pdf_necessite_ocr_texte_insuffisant():
    # 2 pages, quasiment aucun texte -> OCR nécessaire
    segments = [Segment(numero=1, texte="x"), Segment(numero=2, texte="y")]
    assert pdf_necessite_ocr(segments, nb_pages_total=2) is True


def test_pdf_necessite_ocr_texte_suffisant():
    segments = [Segment(numero=1, texte="Un paragraphe bien rempli de texte réel. " * 5)]
    assert pdf_necessite_ocr(segments, nb_pages_total=1) is False


def test_pdf_necessite_ocr_aucune_page():
    assert pdf_necessite_ocr([], nb_pages_total=0) is True


def test_formats_indexables_coherents():
    assert set(FORMATS_INDEXABLES) == {"pdf", "docx", "pptx", "xlsx"}
