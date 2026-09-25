# ============================================================
# TESTS — app/orchestrator/rag_orchestrator.py (Étape C, intégration)
# ============================================================
# Aucune vraie base (faux KnowledgeRepository en mémoire), aucun vrai appel
# Voyage/Groq (faux fournisseur d'embeddings, résumé Groq désactivé sauf
# test dédié), aucun fichier réel du projet touché (registre + dossier des
# fichiers originaux redirigés vers tmp_path).
# ============================================================

import asyncio
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument

import app.orchestrator.rag_orchestrator as rag_orch_module
from app.orchestrator.rag_orchestrator import RagOrchestrator
from app.services.rag import registry_service as registre
from app.services.rag import catalogue_service as catalogue
from app.services.rag import resume_service


def run(coro):
    return asyncio.run(coro)


# ============================================================
# FAUX FOURNISSEUR D'EMBEDDINGS — aucun réseau
# ============================================================

class FauxEmbeddingProvider:
    def __init__(self, tpm=10_000, rpm=3, dimension=1024):
        self.config = SimpleNamespace(tpm=tpm, rpm=rpm)
        self.appels_documents = 0
        self.textes_recus = []
        self._dimension = dimension

    @property
    def modele_document(self):
        return "voyage-4-large"

    @property
    def modele_requete(self):
        return "voyage-4-large"

    async def embed_documents(self, textes):
        self.appels_documents += 1
        self.textes_recus.append(list(textes))
        return [[0.1] * self._dimension for _ in textes]

    async def embed_query(self, texte, *, attente_max_s=5.0):
        return [0.1] * self._dimension


class FauxKnowledgeRepository:
    def __init__(self):
        self.chunks = []
        self.appels_insertion = 0

    def modeles_presents(self, collection=None):
        modeles = {c.modele_embed for c in self.chunks if not collection or c.collection == collection}
        return list(modeles)

    def rechercher_par_similarite(self, vecteur, collection, top_k, **filtres):
        return []  # non utilisé par les tests de l'orchestrateur (Étape D)

    def inserer_chunks(self, lignes):
        self.appels_insertion += 1
        self.chunks.extend(lignes)

    def supprimer_par_hash(self, doc_hash):
        avant = len(self.chunks)
        self.chunks = [c for c in self.chunks if c.doc_hash != doc_hash]
        return avant - len(self.chunks)

    def remplacer_resume_formation(self, formation_code, ligne):
        self.chunks = [
            c for c in self.chunks
            if not (c.collection == "resume_formation" and c.formation_code == formation_code)
        ]
        self.chunks.append(ligne)


@pytest.fixture(autouse=True)
def isole(monkeypatch, tmp_path):
    monkeypatch.setattr(registre, "CHEMIN_REGISTRE_DEFAUT", tmp_path / "index_registry.json")
    monkeypatch.setattr(rag_orch_module, "DOSSIER_FICHIERS_ORIGINAUX", tmp_path / "fichiers")
    # Pas de résumé Groq par défaut (sauf test dédié) : simule un échec propre.
    monkeypatch.setattr(resume_service, "generer_resume_support", _async(lambda *a, **k: None))
    monkeypatch.setattr(resume_service, "generer_resume_formation", _async(lambda *a, **k: None))
    monkeypatch.setattr(rag_orch_module, "generer_resume_support", _async(lambda *a, **k: None))
    monkeypatch.setattr(rag_orch_module, "generer_resume_formation", _async(lambda *a, **k: None))


def _async(fn):
    async def wrapper(*args, **kwargs):
        return fn(*args, **kwargs)
    return wrapper


@pytest.fixture
def orchestrateur(monkeypatch):
    # get_embedding_provider() lirait le vrai .env et ferait de vrais appels
    # réseau : remplacé par un faux AVANT toute construction, par défaut
    # pour tous les tests (certains le réassignent avec un tpm différent).
    monkeypatch.setattr(rag_orch_module, "get_embedding_provider", lambda: FauxEmbeddingProvider())
    orch = RagOrchestrator(repository=FauxKnowledgeRepository())
    return orch


def _docx_simple(tmp_path, nom="doc.docx", texte="Un paragraphe de contenu réel. "):
    chemin = tmp_path / nom
    doc = DocxDocument()
    doc.add_paragraph(texte)
    doc.save(str(chemin))
    return chemin


# ============================================================
# CLASSIFICATION
# ============================================================

def test_ingerer_fichier_image_seule_non_indexable(orchestrateur, tmp_path):
    chemin = tmp_path / "photo.jpg"
    chemin.write_bytes(bytes.fromhex("FFD8FFE000104A46494600010100000100010000FFD9"))

    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert resultat["statut"] == "non_indexable_image"
    assert orchestrateur.embedding_provider.appels_documents == 0


def test_ingerer_fichier_format_video_refuse(orchestrateur, tmp_path):
    chemin = tmp_path / "video.mp4"
    chemin.write_bytes(b"contenu binaire quelconque")

    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert resultat["statut"] == "erreur"
    assert "vidéo" in resultat["message"].lower() or "audio" in resultat["message"].lower()


def test_ingerer_fichier_pdf_scanne_ocr_necessaire(orchestrateur, tmp_path, monkeypatch):
    chemin = tmp_path / "scan.pdf"
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(chemin))
    c.showPage()  # page vide, aucun texte
    c.save()

    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert resultat["statut"] == "ocr_necessaire"
    assert orchestrateur.embedding_provider.appels_documents == 0


# ============================================================
# CONFIDENTIALITÉ
# ============================================================

def test_ingerer_fichier_formation_confidentielle_refusee(orchestrateur, tmp_path, monkeypatch):
    chemin_catalogue = tmp_path / "catalogue.yaml"
    chemin_catalogue.write_text(
        "formations:\n  - code: SECRET\n    titre: 'Formation secrète'\n    confidentiel: true\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(catalogue, "CHEMIN_CATALOGUE_DEFAUT", chemin_catalogue)
    monkeypatch.setattr(rag_orch_module, "obtenir_formation", lambda code: catalogue.obtenir_formation(code, chemin_catalogue))
    monkeypatch.delenv("VOYAGE_OPT_OUT", raising=False)

    chemin_doc = _docx_simple(tmp_path)
    resultat = run(orchestrateur.ingerer_fichier(str(chemin_doc), formation_code="SECRET"))

    assert resultat["statut"] == "erreur"
    assert "confidentielle" in resultat["message"].lower()
    assert orchestrateur.embedding_provider.appels_documents == 0


# ============================================================
# IDEMPOTENCE
# ============================================================

def test_ingerer_fichier_idempotence(orchestrateur, tmp_path):
    chemin = _docx_simple(tmp_path)

    premier = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))
    appels_apres_premier = orchestrateur.embedding_provider.appels_documents

    second = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert premier["statut"] == "indexe"
    assert second["statut"] == "deja_indexe"
    assert orchestrateur.embedding_provider.appels_documents == appels_apres_premier  # aucun nouvel appel Voyage


# ============================================================
# ÉCHEC GROQ — document indexé quand même
# ============================================================

def test_echec_groq_document_indexe_quand_meme(orchestrateur, tmp_path):
    chemin = _docx_simple(tmp_path)

    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert resultat["statut"] == "indexe"  # le résumé Groq est mocké à None (échec simulé)
    entree = registre.obtenir_entree(resultat["hash"])
    assert entree.resume_statut == "a_generer"


# ============================================================
# REPRISE APRÈS INTERRUPTION
# ============================================================

def test_reprise_apres_interruption_ne_refait_pas_les_lots_termines(orchestrateur, tmp_path):
    # TPM très bas : chaque chunk (~700 tokens) dépasse déjà le budget d'un
    # lot -> 1 chunk = 1 lot, donc plusieurs lots pour un document à
    # plusieurs chunks (déterministe, sans dépendre du texte exact).
    orchestrateur.embedding_provider = FauxEmbeddingProvider(tpm=50)

    chemin = tmp_path / "long.docx"
    doc = DocxDocument()
    for i in range(20):
        doc.add_paragraph(f"Paragraphe numéro {i}. " + ("Contenu substantiel pour remplir le chunk. " * 40))
    doc.save(str(chemin))

    hash_fichier = registre.calculer_hash_fichier(str(chemin))
    registre.initialiser_entree(hash_fichier, "long.docx", None)
    registre.marquer_lot_termine(hash_fichier, 0)  # simule : lot 0 déjà inséré lors d'une exécution précédente

    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))

    assert resultat["statut"] == "indexe"
    # Le lot 0 n'a PAS été redemandé à Voyage : le nombre d'appels est
    # strictement inférieur au nombre total de chunks (1 lot = 1 chunk ici).
    assert orchestrateur.embedding_provider.appels_documents < resultat["nb_chunks"]


# ============================================================
# SUPPRESSION
# ============================================================

def test_supprimer_document_retire_les_chunks_et_le_registre(orchestrateur, tmp_path):
    chemin = _docx_simple(tmp_path)
    resultat = run(orchestrateur.ingerer_fichier(str(chemin), formation_code=None))
    assert len(orchestrateur.repository.chunks) > 0

    suppression = run(orchestrateur.supprimer_document(resultat["hash"]))

    assert suppression["statut"] == "supprime"
    assert len(orchestrateur.repository.chunks) == 0
    assert registre.obtenir_entree(resultat["hash"]) is None


def test_supprimer_document_introuvable(orchestrateur):
    resultat = run(orchestrateur.supprimer_document("hash-jamais-vu"))
    assert resultat["statut"] == "introuvable"


# ============================================================
# ESTIMATION DU COÛT
# ============================================================

def test_estimer_cout(orchestrateur):
    orchestrateur.embedding_provider = FauxEmbeddingProvider(tpm=1000)
    textes = ["x" * 400] * 10  # 10 * 100 tokens = 1000 tokens

    estimation = orchestrateur.estimer_cout(textes)

    assert estimation["tokens_estimes"] == 1000
    assert estimation["nb_requetes"] >= 1
    assert "duree_estimee_min" in estimation


# ============================================================
# ZIP — rapport par fichier, un échec n'arrête pas les autres
# ============================================================

def test_ingerer_zip_rapport_par_fichier(orchestrateur, tmp_path):
    chemin_docx = tmp_path / "bon.docx"
    doc = DocxDocument()
    doc.add_paragraph("Contenu valide et suffisamment long pour former un chunk exploitable.")
    doc.save(str(chemin_docx))

    chemin_zip = tmp_path / "corpus.zip"
    with zipfile.ZipFile(chemin_zip, "w") as zf:
        zf.write(chemin_docx, arcname="bon.docx")
        zf.writestr("image.jpg", bytes.fromhex("FFD8FFE000104A46494600010100000100010000FFD9"))
        zf.writestr("video.mp4", b"contenu video")

    resultat = run(orchestrateur.ingerer_zip(str(chemin_zip), formation_code=None))

    assert resultat["statut"] == "termine"
    assert resultat["nb_fichiers"] == 3
    assert resultat["nb_indexes"] == 1
    statuts = {r["fichier"]: r["statut"] for r in resultat["fichiers"]}
    assert statuts["bon.docx"] == "indexe"
    assert statuts["image.jpg"] == "non_indexable_image"
    assert statuts["video.mp4"] == "erreur"


def test_ingerer_zip_bombe_refuse_avant_extraction(orchestrateur, tmp_path, monkeypatch):
    from app.services.rag import zip_securite
    monkeypatch.setattr(zip_securite, "MAX_DECOMPRESSE_OCTETS", 10)

    chemin_zip = tmp_path / "bombe.zip"
    with zipfile.ZipFile(chemin_zip, "w") as zf:
        zf.writestr("gros.txt", b"x" * 1000)

    resultat = run(orchestrateur.ingerer_zip(str(chemin_zip), formation_code=None))

    assert resultat["statut"] == "refuse"
    assert orchestrateur.embedding_provider.appels_documents == 0
